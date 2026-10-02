"""
actions/file_processor.py — Tahtaya bırakılan belge/resim işleyici.

Dosya doğrudan Gemini'ye gider (`core/saglayicilar.py`), başka bir sunucu
yoktur:
  - PDF ve resimler ikili içerik olarak gönderilir — Gemini sayfa düzenini,
    tabloları ve el yazısını kendisi okur.
  - Metin tabanlı dosyalar (txt, md, csv, json, xml, html) metin olarak
    istemin içine konur.
  - Word (.docx) ve Excel (.xlsx/.xlsm): Gemini bunları doğrudan okumaz;
    metin yerelde çıkarılır (python-docx: paragraflar + tablolar; openpyxl:
    her sayfa, hücre değerleri) ve metin olarak gönderilir.
  - PowerPoint ve eski ikili biçimler (.doc/.xls/.ppt) ile OpenDocument
    desteklenmez; kullanıcıya dosyayı PDF olarak kaydetmesi söylenir.

Dosyaya geri yazma, biçim dönüştürme ve kaydetme yok; araç yalnızca okur ve
modele metin bir sonuç döndürür.
"""

from pathlib import Path

from core import saglayicilar

AZAMI_BOYUT = 15 * 1024 * 1024          # satır içi istek sınırının altında kalsın
AZAMI_METIN = 200_000                   # metin dosyalarında karakter sınırı

_RESIM = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
          ".webp": "image/webp", ".heic": "image/heic", ".heif": "image/heif"}
_METIN = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm"}
_WORD = {".docx"}
_EXCEL = {".xlsx", ".xlsm"}
_OFIS = {".doc", ".xls", ".ppt", ".pptx", ".odt", ".ods", ".odp"}

_EYLEM_ISTEMI = {
    "summarize":    "Bu belgeyi lise düzeyinde, Türkçe ve kısa maddelerle özetle.",
    "extract_text": "Bu belgedeki metni olduğu gibi çıkar; yorum ekleme.",
    "ocr":          "Bu görseldeki tüm yazıyı (formüller dahil) olduğu gibi çıkar; yorum ekleme.",
    "describe":     "Bu görseli bir öğretmene anlatır gibi Türkçe betimle.",
    "analyze":      "Bu veriyi incele; öne çıkan bulguları Türkçe ve kısa maddelerle yaz.",
    "stats":        "Bu verinin temel istatistiklerini (satır/sütun sayısı, sayısal sütunların ortalama/en küçük/en büyük değerleri) çıkar.",
}

_SISTEM = (
    "Bir sınıf asistanı için dosya okuyorsun. Yalnızca dosyada gerçekten olanı "
    "söyle; dosyada olmayan bilgi, sayı ya da kaynak uydurma. Okunamayan bir "
    "kısım varsa bunu açıkça belirt."
)


def _word_metni(path: Path) -> str:
    """Paragraflar sırayla, ardından tablolar (hücreler ' | ' ile)."""
    import docx

    belge = docx.Document(str(path))
    satirlar = [p.text for p in belge.paragraphs if p.text.strip()]
    for no, tablo in enumerate(belge.tables, 1):
        satirlar.append(f"\n[Tablo {no}]")
        for satir in tablo.rows:
            satirlar.append(" | ".join(h.text.strip() for h in satir.cells))
    return "\n".join(satirlar)


def _excel_metni(path: Path) -> str:
    """Her sayfa başlığıyla, satırlar sekmeyle ayrılmış değerler olarak.
    Formüllerin HESAPLANMIŞ değeri okunur (data_only); boş satırlar atlanır."""
    import openpyxl

    kitap = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        parcalar = []
        for sayfa in kitap.worksheets:
            parcalar.append(f"[Sayfa: {sayfa.title}]")
            for satir in sayfa.iter_rows(values_only=True):
                if any(h is not None and str(h).strip() for h in satir):
                    parcalar.append("\t".join("" if h is None else str(h) for h in satir))
            if sum(len(p) for p in parcalar) > AZAMI_METIN:
                break
        return "\n".join(parcalar)
    finally:
        kitap.close()


def _istem(eylem: str, talimat: str, varsayilan: str) -> str:
    temel = _EYLEM_ISTEMI.get(eylem, varsayilan)
    return f"{temel}\nEk talimat: {talimat}" if talimat else temel


def file_processor(parameters: dict, player=None, speak=None) -> str:
    log = getattr(player, "write_log", None) or (lambda *_a: None)

    file_path_str = (parameters.get("file_path") or "").strip()
    if not file_path_str:
        return "Dosya yolu belirtilmedi."

    path = Path(file_path_str)
    if not path.exists():
        return f"Dosya bulunamadı: {file_path_str}"
    if not path.is_file():
        return f"Yol bir dosya değil: {file_path_str}"
    if path.stat().st_size > AZAMI_BOYUT:
        return "Dosya çok büyük (en fazla 15 MB). Daha küçük bir dosya ya da birkaç sayfa deneyin."

    eylem = (parameters.get("action") or "").lower().strip()
    talimat = (parameters.get("instruction") or "").strip()
    uzanti = path.suffix.lower()
    log(f"[FileProcessor] {path.name} | action={eylem or 'auto'}")

    try:
        if uzanti == ".pdf":
            return saglayicilar.gorsel_uret(
                "belge_ozet", _istem(eylem, talimat, _EYLEM_ISTEMI["summarize"]),
                path.read_bytes(), "application/pdf", sistem=_SISTEM)
        if uzanti in _RESIM:
            return saglayicilar.gorsel_uret(
                "gorsel_okuma", _istem(eylem, talimat, _EYLEM_ISTEMI["describe"]),
                path.read_bytes(), _RESIM[uzanti], sistem=_SISTEM)
        if uzanti in _METIN or uzanti in _WORD or uzanti in _EXCEL:
            if uzanti in _WORD:
                metin, varsayilan = _word_metni(path), _EYLEM_ISTEMI["summarize"]
            elif uzanti in _EXCEL:
                metin, varsayilan = _excel_metni(path), _EYLEM_ISTEMI["analyze"]
            else:
                metin = path.read_text(encoding="utf-8", errors="replace")
                varsayilan = _EYLEM_ISTEMI["summarize"]
            if not metin.strip():
                return "Dosyada okunabilir metin bulunamadı."
            kirpildi = len(metin) > AZAMI_METIN
            metin = metin[:AZAMI_METIN]
            istem = (f"{_istem(eylem, talimat, varsayilan)}\n\n"
                     f"--- {path.name}{' (ilk kısmı, dosya kırpıldı)' if kirpildi else ''} ---\n{metin}")
            return saglayicilar.metin_uret("belge_ozet", istem, sistem=_SISTEM)
    except RuntimeError as e:
        log(f"[FileProcessor] hata: {e}")
        return ("Dosya şu an okunamadı. Dosya içeriği hakkında tahmin yürütme; "
                "öğretmene dosyayı daha sonra tekrar denemesini söyle.")
    except Exception as e:      # bozuk/şifreli docx-xlsx, eksik kütüphane...
        log(f"[FileProcessor] {path.name} açılamadı: {type(e).__name__}: {e}")
        return ("Dosya açılamadı (bozuk ya da şifreli olabilir). İçeriği hakkında "
                "tahmin yürütme; öğretmenden dosyayı PDF olarak kaydetmesini iste.")

    if uzanti in _OFIS:
        return ("Bu dosya türü (PowerPoint, eski .doc/.xls ya da OpenDocument) "
                "okunamıyor. Öğretmenden dosyayı PDF olarak kaydedip tekrar "
                "bırakmasını iste.")
    return f"Desteklenmeyen dosya türü: {uzanti or 'uzantısız'}."
